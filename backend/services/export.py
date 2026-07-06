import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx

def format_text_data(data: dict) -> str:
    """Formats summaries, quizzes, flashcards, and key points into a single raw text string."""
    txt = []
    txt.append(f"AI RESEARCH ASSISTANT REPORT - {data.get('title', 'Document Report')}\n")
    txt.append("="*80 + "\n\n")
    
    # Summary
    if "summary" in data:
        sum_data = data["summary"]
        txt.append("1. DOCUMENT SUMMARY\n")
        txt.append("-" * 30 + "\n")
        txt.append(f"Short Summary:\n{sum_data.get('short', '')}\n\n")
        txt.append(f"Medium Summary:\n{sum_data.get('medium', '')}\n\n")
        txt.append(f"Detailed Summary:\n{sum_data.get('detailed', '')}\n\n")
        
    # Key Points
    if "key_points" in data and data["key_points"]:
        txt.append("2. KEY POINTS & FINDINGS\n")
        txt.append("-" * 30 + "\n")
        for kp in data["key_points"]:
            txt.append(f"• {kp}\n")
        txt.append("\n")
        
    # Quiz
    if "quiz" in data and data["quiz"]:
        txt.append("3. COMPREHENSION QUIZ\n")
        txt.append("-" * 30 + "\n")
        for idx, q in enumerate(data["quiz"]):
            txt.append(f"Q{idx+1}: {q.get('question', '')}\n")
            for i, opt in enumerate(q.get('options', [])):
                txt.append(f"  {chr(65+i)}) {opt}\n")
            txt.append(f"Correct Answer: {q.get('correct_answer', '')}\n")
            txt.append(f"Explanation: {q.get('explanation', '')}\n\n")
            
    # Flashcards
    if "flashcards" in data and data["flashcards"]:
        txt.append("4. FLASHCARDS FOR STUDY\n")
        txt.append("-" * 30 + "\n")
        for idx, fc in enumerate(data["flashcards"]):
            txt.append(f"Card {idx+1}:\n")
            txt.append(f"  Front: {fc.get('front', '')}\n")
            txt.append(f"  Back:  {fc.get('back', '')}\n\n")
            
    return "".join(txt)

def export_to_docx(data: dict, file_path: str):
    """Generates a formatted Word Document (.docx) using python-docx."""
    doc = docx.Document()
    
    # Title
    doc.add_heading(data.get('title', 'AI Research Agent Export'), 0)
    
    # Summary
    if "summary" in data:
        doc.add_heading('1. Document Summaries', level=1)
        sum_data = data["summary"]
        doc.add_heading('Short Summary', level=2)
        doc.add_paragraph(sum_data.get('short', ''))
        
        doc.add_heading('Medium Summary', level=2)
        doc.add_paragraph(sum_data.get('medium', ''))
        
        doc.add_heading('Detailed Summary', level=2)
        doc.add_paragraph(sum_data.get('detailed', ''))
        
    # Key Points
    if "key_points" in data and data["key_points"]:
        doc.add_heading('2. Key Points & Findings', level=1)
        for kp in data["key_points"]:
            doc.add_paragraph(kp, style='List Bullet')
            
    # Quiz
    if "quiz" in data and data["quiz"]:
        doc.add_heading('3. Comprehension Quiz', level=1)
        for idx, q in enumerate(data["quiz"]):
            doc.add_heading(f"Q{idx+1}: {q.get('question', '')}", level=2)
            for i, opt in enumerate(q.get('options', [])):
                doc.add_paragraph(f"{chr(65+i)}) {opt}")
            doc.add_paragraph(f"Correct Answer: {q.get('correct_answer', '')}").bold = True
            doc.add_paragraph(f"Explanation: {q.get('explanation', '')}")
            
    # Flashcards
    if "flashcards" in data and data["flashcards"]:
        doc.add_heading('4. Study Flashcards', level=1)
        for idx, fc in enumerate(data["flashcards"]):
            doc.add_heading(f"Flashcard {idx+1}", level=2)
            doc.add_paragraph(f"Front (Concept/Question): {fc.get('front', '')}")
            doc.add_paragraph(f"Back (Definition/Answer): {fc.get('back', '')}")
            
    doc.save(file_path)

def export_to_pdf(data: dict, file_path: str):
    """Generates a structured PDF Document using ReportLab."""
    doc = SimpleDocTemplate(file_path, pagesize=letter)
    styles = getSampleStyleSheet()
    
    # Create custom styles
    title_style = ParagraphStyle(
        'DocTitle', 
        parent=styles['Heading1'], 
        fontSize=22, 
        spaceAfter=15, 
        textColor=colors.HexColor('#4F46E5') # Indigo
    )
    h1_style = ParagraphStyle(
        'H1Style', 
        parent=styles['Heading2'], 
        fontSize=15, 
        spaceBefore=12, 
        spaceAfter=6, 
        textColor=colors.HexColor('#0F172A') # Slate 900
    )
    h2_style = ParagraphStyle(
        'H2Style', 
        parent=styles['Heading3'], 
        fontSize=11, 
        spaceBefore=8, 
        spaceAfter=4, 
        textColor=colors.HexColor('#1E293B')
    )
    body_style = ParagraphStyle(
        'BodyStyle', 
        parent=styles['Normal'], 
        fontSize=10, 
        leading=14, 
        spaceAfter=8
    )
    bullet_style = ParagraphStyle(
        'BulletStyle', 
        parent=styles['Normal'], 
        fontSize=10, 
        leading=14, 
        leftIndent=15, 
        spaceAfter=4
    )
    bold_body_style = ParagraphStyle(
        'BoldBodyStyle', 
        parent=body_style, 
        fontName='Helvetica-Bold'
    )
    
    story = []
    
    # Title
    story.append(Paragraph(data.get('title', 'AI Research Agent Export'), title_style))
    story.append(Spacer(1, 10))
    
    # Summary
    if "summary" in data:
        story.append(Paragraph('1. Document Summaries', h1_style))
        sum_data = data["summary"]
        
        story.append(Paragraph('Short Summary', h2_style))
        story.append(Paragraph(sum_data.get('short', ''), body_style))
        
        story.append(Paragraph('Medium Summary', h2_style))
        story.append(Paragraph(sum_data.get('medium', ''), body_style))
        
        story.append(Paragraph('Detailed Summary', h2_style))
        story.append(Paragraph(sum_data.get('detailed', ''), body_style))
        story.append(Spacer(1, 10))
        
    # Key Points
    if "key_points" in data and data["key_points"]:
        story.append(Paragraph('2. Key Points & Findings', h1_style))
        for kp in data["key_points"]:
            story.append(Paragraph(f"• {kp}", bullet_style))
        story.append(Spacer(1, 10))
        
    # Quiz
    if "quiz" in data and data["quiz"]:
        story.append(Paragraph('3. Comprehension Quiz', h1_style))
        for idx, q in enumerate(data["quiz"]):
            story.append(Paragraph(f"Q{idx+1}: {q.get('question', '')}", h2_style))
            for i, opt in enumerate(q.get('options', [])):
                story.append(Paragraph(f"{chr(65+i)}) {opt}", bullet_style))
            story.append(Paragraph(f"Correct Answer: {q.get('correct_answer', '')}", bold_body_style))
            story.append(Paragraph(f"Explanation: {q.get('explanation', '')}", body_style))
            story.append(Spacer(1, 5))
        story.append(Spacer(1, 10))
            
    # Flashcards
    if "flashcards" in data and data["flashcards"]:
        story.append(Paragraph('4. Study Flashcards', h1_style))
        for idx, fc in enumerate(data["flashcards"]):
            story.append(Paragraph(f"Flashcard {idx+1}", h2_style))
            story.append(Paragraph(f"Front: {fc.get('front', '')}", body_style))
            story.append(Paragraph(f"Back: {fc.get('back', '')}", body_style))
            story.append(Spacer(1, 5))
            
    doc.build(story)
